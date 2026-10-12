import os
import io
import re
import datetime
import tempfile
import shutil
from pathlib import Path
from starlette.background import BackgroundTask
from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Response, Request
from fastapi.responses import StreamingResponse, FileResponse
from sqlalchemy.orm import Session
import openpyxl

from backend.app.database import get_db, get_business_tables, BASE_DIR, DB_PATH, DATA_DIR, engine
from backend.app.services.database_backup import snapshot, restore_database
from backend.app.models.models import Mouse, Cage, Claimer, TransferLog, Primer, User
from backend.app.services.importer import import_local_excel_folder, import_single_excel_file
from backend.app.services.owner_service import sync_euthanasia_owner
from backend.app.services.strain_service import normalize_strain_name
from backend.app.auth import require_admin, require_auth

router = APIRouter(prefix="/api/import-export", tags=["Import & Export"])


@router.get('/database/export')
def export_database(current_user: User = Depends(require_admin)):
    folder = tempfile.mkdtemp(prefix='mouse-backup-')
    path = os.path.join(folder, 'backup.db')
    try:
        snapshot(DB_PATH, path)
    except Exception:
        shutil.rmtree(folder)
        raise
    filename = f"mouse-manager_{datetime.datetime.now():%Y%m%d_%H%M%S}.db"
    return FileResponse(path, filename=filename, media_type='application/octet-stream',
                        background=BackgroundTask(shutil.rmtree, folder))


@router.post('/database/restore')
def import_database(file: UploadFile = File(...), db: Session = Depends(get_db),
                    current_user: User = Depends(require_admin)):
    if not (file.filename or '').lower().endswith('.db'):
        raise HTTPException(status_code=400, detail='请选择 .db 数据库备份文件')
    with tempfile.TemporaryDirectory(prefix='mouse-restore-') as folder:
        path = Path(folder) / 'restore.db'
        with path.open('wb') as output:
            size = 0
            while chunk := file.file.read(1024 * 1024):
                size += len(chunk)
                if size > 512 * 1024 * 1024:
                    raise HTTPException(status_code=400, detail='数据库文件不能超过 512 MB')
                output.write(chunk)
        db.close()
        backup_dir = Path(DATA_DIR) / 'backups'
        backup_dir.mkdir(exist_ok=True)
        safety = backup_dir / f'before_restore_{datetime.datetime.now():%Y%m%d_%H%M%S_%f}.db'
        try:
            restore_database(path, DB_PATH, safety, get_business_tables(), engine)
            with Session(engine) as restored_db:
                sync_euthanasia_owner(restored_db)
                restored_db.commit()
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {'message': '小鼠业务数据库已完整恢复，账号数据库保持不变',
            'safety_backup': safety.name}

def style_export_header(sheet):
    for cell in sheet[1]:
        cell.font = openpyxl.styles.Font(bold=True)
        cell.fill = openpyxl.styles.PatternFill(fill_type="solid", fgColor="FFE7E6E6")

@router.post("/init-local")
def init_from_local_excel(db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    """One-click import from workspace excel/ folder"""
    excel_dir = os.path.join(BASE_DIR, "excel")
    if not os.path.exists(excel_dir):
        raise HTTPException(status_code=404, detail="未找到 excel 文件夹")

    results = import_local_excel_folder(db, excel_dir, operator=current_user.display_name or current_user.username)
    return {
        "success": True,
        "message": "导入完成",
        "details": results
    }

@router.post("/upload")
async def upload_excel(
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin)
):
    form = await request.form()
    upload_files = [
        v for k, v in form.multi_items()
        if hasattr(v, "filename") and getattr(v, "filename", None)
    ]
    if not upload_files:
        raise HTTPException(status_code=400, detail="请选择要上传的 Excel 文件 (.xlsx 或 .xls)")

    temp_dir = os.path.join(BASE_DIR, "data", "uploads")
    os.makedirs(temp_dir, exist_ok=True)

    combined_results = {
        "mice_imported": 0,
        "cages_imported": 0,
        "claimers_imported": 0,
        "transfers_imported": 0,
        "transfer_requests_imported": 0,
        "genotypes_imported": 0,
        "inferred_genders_count": 0,
        "primers_imported": 0,
        "errors": [],
        "report": {
            "summary": {"added": 0, "updated": 0, "skipped": 0, "removed": 0},
            "records": [], "ignored_sheets": [],
        },
    }
    file_summaries = []

    for i, file in enumerate(upload_files):
        if not file.filename.endswith((".xlsx", ".xls")):
            err = "仅支持上传 .xlsx 或 .xls 格式文件"
            combined_results["errors"].append(f"{file.filename}: {err}")
            file_summaries.append({
                "filename": file.filename,
                "success": False,
                "error": err,
                "imported_count": 0
            })
            continue

        temp_path = os.path.join(
            temp_dir,
            f"{datetime.datetime.now().strftime('%Y%m%d%H%M%S')}_{i}_{file.filename}"
        )
        content = await file.read()
        with open(temp_path, "wb") as f:
            f.write(content)

        try:
            results = import_single_excel_file(
                db, temp_path, original_filename=file.filename,
                operator=current_user.display_name or current_user.username,
            )
            for k in [
                "mice_imported", "cages_imported", "claimers_imported",
                "transfers_imported", "transfer_requests_imported",
                "genotypes_imported", "inferred_genders_count", "primers_imported"
            ]:
                combined_results[k] += results.get(k, 0)
            if results.get("errors"):
                combined_results["errors"].extend(results["errors"])

            report = results["report"]
            for status, count in report["summary"].items():
                combined_results["report"]["summary"][status] += count
            combined_results["report"]["records"].extend(report["records"])
            combined_results["report"]["ignored_sheets"].extend(report["ignored_sheets"])

            file_items = report["summary"]["added"] + report["summary"]["updated"]
            file_summaries.append({
                "filename": file.filename,
                "success": not results["errors"] or file_items > 0,
                "imported_count": file_items,
                "error": "；".join(results["errors"]),
                "details": results
            })
        except Exception as e:
            combined_results["errors"].append(f"{file.filename} 解析失败: {str(e)}")
            file_summaries.append({
                "filename": file.filename,
                "success": False,
                "error": str(e),
                "imported_count": 0
            })

    await form.close()
    summary = combined_results["report"]["summary"]
    total_items = summary["added"] + summary["updated"]
    summary_msg = f"导入完成：新增 {summary['added']} 项，补充/更新 {summary['updated']} 项，重复跳过 {summary['skipped']} 项"
    if summary["removed"]:
        summary_msg += f"，清理无效记录 {summary['removed']} 项"
    if combined_results["errors"]:
        summary_msg += f"，解析失败 {len(combined_results['errors'])} 项（详见明细）"

    return {
        "success": True,
        "file_count": len(upload_files),
        "imported_type": "Excel工作表综合解析",
        "imported_count": total_items,
        "details": combined_results,
        "file_summaries": file_summaries,
        "message": summary_msg
    }

@router.get("/export/mice")
def export_mice(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_auth)
):
    """Export mice to separate worksheets by room."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    headers = [
        "耳标编号", "品系/基因型", "性别", "出生日期", "周龄", 
        "所在鼠房", "所在笼位", "领取人", "领用日期", "领用目的",
        "Genotype 1", "Genotype 2", "测试日期", "父母来源", "状态", "备注"
    ]
    room_sheets = {}

    def get_room_sheet(room):
        if room not in room_sheets:
            # Excel limits titles to 31 characters and forbids these characters.
            title = re.sub(r'[\\/*?:\[\]\x00-\x1f]', '_', room or '未分配鼠房').strip("'")
            title = title[:31] or '未分配鼠房'
            used_titles = {name.casefold() for name in wb.sheetnames}
            candidate = title
            suffix_number = 2
            while candidate.casefold() in used_titles:
                suffix = f' ({suffix_number})'
                candidate = title[:31 - len(suffix)] + suffix
                suffix_number += 1
            sheet = wb.create_sheet(candidate)
            sheet.append(headers)
            style_export_header(sheet)
            room_sheets[room] = sheet
        return room_sheets[room]

    mice = db.query(Mouse).order_by(Mouse.id.desc()).all()
    for m in mice:
        room = ((m.cage.room if m.cage else None) or m.source_room or '').strip()
        ws = get_room_sheet(room)
        row = ws.max_row + 1
        dob_cell = f'D{row}'
        age_weeks = (
            f'=IF({dob_cell}="","",IFERROR(ROUND((TODAY()-'
            f'IF(ISNUMBER({dob_cell}),{dob_cell},DATEVALUE({dob_cell})))/7,1),""))'
        )

        ws.append([
            m.mouse_code,
            m.strain or "",
            m.gender or "",
            m.dob or "",
            age_weeks,
            room,
            m.cage.cage_code if m.cage else "",
            m.owner_name or "",
            m.claim_date or "",
            m.claim_purpose or "",
            m.genotype_1 or "",
            m.genotype_2 or "",
            m.test_date or "",
            m.parents or "",
            m.status or "",
            m.notes or ""
        ])

    if not room_sheets:
        get_room_sheet('')
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"mice_export_{datetime.date.today().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@router.get("/export/cages")
def export_cages(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_auth)
):
    """Export cages to Excel"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "小鼠笼位导出"

    headers = ["鼠房", "笼位号", "品系", "性别", "容量", "在笼数量", "在笼小鼠耳标", "合笼日期", "生鼠日期", "观察记录", "备注"]
    ws.append(headers)
    style_export_header(ws)

    cages = db.query(Cage).order_by(Cage.room.asc(), Cage.cage_code.asc()).all()
    for c in cages:
        mice_codes = ", ".join([m.mouse_code for m in c.mice])
        ws.append([
            c.room,
            c.cage_code,
            c.strain or "",
            c.gender or "",
            c.capacity,
            len(c.mice),
            mice_codes,
            c.mating_date or "",
            "、".join(c.litter_birth_dates or ([c.litter_birth_date] if c.litter_birth_date else [])),
            c.observation or "",
            c.notes or ""
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"cages_export_{datetime.date.today().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
