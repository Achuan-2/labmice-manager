import unittest
from itertools import permutations
from unittest.mock import patch

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from backend.app.database import Base
from backend.app.models.models import GenotypeRecord, Mouse, Strain, User
from backend.app.routers.genotypes import list_genotypes
from backend.app.services.importer import import_single_excel_file, import_workbook
from backend.app.services.strain_service import backfill_known_pedigree_strains, infer_strain_from_genotypes, sync_and_normalize_all_strains


class GenotypeStrainTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def import_rows(self, *rows):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "基因鉴定结果"
        sheet.append(["测试日期", "小鼠编号", "品系", "DOB", "性别", "父母", "Genotype 1", "Genotype 2", "Genotype 3"])
        for row in rows:
            sheet.append(row)
        results = {key: 0 for key in [
            "mice_imported", "cages_imported", "claimers_imported", "transfers_imported",
            "transfer_requests_imported", "genotypes_imported", "inferred_genders_count", "primers_imported",
        ]}
        results["errors"] = []
        import_workbook(self.db, workbook, results, {}, "test.xlsx")
        self.assertEqual(results["errors"], [])

    def test_explicit_strain_is_not_validated_as_an_ear_tag(self):
        self.import_rows(
            ["2026-08-18", "F335", "Camk2/Ai93-Ras-N", "2026-07-05", "F", None, "Camk2a-阳性", "Ai93-杂合子", "RasN-杂合子"],
            ["2026-08-19", "F340", "自定义品系", None, "M", None, "WT", None, None],
        )
        for code, strain in [("F335", "Camk2/Ai93-Ras-N"), ("F340", "自定义品系")]:
            self.assertEqual(self.db.query(Mouse).filter_by(mouse_code=code).one().strain, strain)
            self.assertEqual(self.db.query(GenotypeRecord).filter_by(mouse_code=code).one().strain, strain)

    def test_reimport_fills_existing_genotype_strain_without_duplicate_or_overwrite(self):
        mouse = Mouse(mouse_code="B829", strain="", genotype_1="WT")
        record = GenotypeRecord(mouse=mouse, mouse_code="B829", strain="", test_date="2023-02-07", genotype_1="WT")
        explicit = GenotypeRecord(mouse_code="B830", strain="已有品系", test_date="2023-02-07", genotype_1="HET")
        self.db.add_all([record, explicit])
        self.db.commit()
        self.import_rows(
            ["2023-02-07", "B829", "JAX-5XFAD-J", "2023-01-10", "M", None, "WT", None, None],
            ["2023-02-07", "B830", "JAX-5XFAD-J", "2023-01-10", "M", None, "HET", None, None],
        )
        self.assertEqual((mouse.strain, record.strain), ("JAX-5XFAD-J", "JAX-5XFAD-J"))
        self.assertEqual(record.genotype_1, "WT")
        self.assertEqual(explicit.strain, "已有品系")
        self.assertEqual(self.db.query(GenotypeRecord).count(), 2)

    def test_import_reads_vertical_strain_merges_across_different_dates(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "小鼠基因型鉴定结果"
        sheet.append(["日期", "编号", "品系", "生日", "Genotype 1"])
        sheet.append(["2023-02-07", "B829", "JAX-5XFAD-J", "2023-01-10", "WT"])
        sheet.append(["2023-02-08", "B830", None, None, "HET"])
        sheet.merge_cells("C2:C3")
        sheet.merge_cells("D2:D3")
        with patch("backend.app.services.importer.openpyxl.load_workbook", return_value=workbook):
            result = import_single_excel_file(self.db, "test.xlsx")
        self.assertEqual(result["errors"], [])
        record = self.db.query(GenotypeRecord).filter_by(mouse_code="B830").one()
        self.assertEqual((record.strain, record.dob, record.test_date), ("JAX-5XFAD-J", "2023-01-10", "2023-02-08"))

    def test_same_date_inherits_merged_strain_but_new_date_resets_it(self):
        self.import_rows(
            ["2026-08-18", "F335", "Camk2/Ai93-Ras-N", "2026-07-05", "F", None, "WT", None, None],
            ["2026-08-18", "F336", None, None, "F", None, "HET", None, None],
            ["2026-08-19", "F337", None, None, "F", None, "HET", None, None],
        )
        same_date = self.db.query(GenotypeRecord).filter_by(mouse_code="F336").one()
        self.assertEqual((same_date.strain, same_date.dob), ("Camk2/Ai93-Ras-N", "2026-07-05"))
        new_date = self.db.query(GenotypeRecord).filter_by(mouse_code="F337").one()
        self.assertEqual(new_date.strain, "")
        self.assertIsNone(new_date.dob)

    def test_missing_strain_uses_only_the_three_named_markers(self):
        self.import_rows(
            ["2026-08-18", "F335", None, None, "F", None, "Camk2a-阳性", "Ai93-杂合子", "RasN-杂合子"],
            ["2026-08-18", "F336", None, None, "F", None, "阳性", "纯合子", "野生型"],
        )
        self.assertEqual(self.db.query(Mouse).filter_by(mouse_code="F335").one().strain, "Camk2/Ai93-Ras-N")
        self.assertEqual(self.db.query(Mouse).filter_by(mouse_code="F336").one().strain, "")

    def test_marker_recognition_does_not_require_positive_outcomes(self):
        self.assertEqual(infer_strain_from_genotypes("Camk2a-阴性", "Ai93-纯合子", "RasN-野生型"), "Camk2/Ai93-Ras-N")
        self.assertEqual(infer_strain_from_genotypes("camk2a-tTA: HET", "ai93: HOM", "Ras-N: WT"), "Camk2/Ai93-Ras-N")
        self.assertEqual(infer_strain_from_genotypes("Camk2a-阳性", "Ai93-纯合子", "Other-野生型"), "")
        self.assertEqual(infer_strain_from_genotypes("Camk2a-阳性", "Ai93-纯合子", None), "")

    def test_normalization_repairs_legacy_data_and_list_response(self):
        mouse = Mouse(mouse_code="F335", strain="")
        record = GenotypeRecord(
            mouse=mouse, mouse_code="F335", strain="", test_date="2026-08-18",
            genotype_1="Camk2a-阳性", genotype_2="Ai93-杂合子", genotype_3="RasN-杂合子",
        )
        self.db.add(record)
        self.db.commit()
        result = sync_and_normalize_all_strains(self.db)
        self.assertEqual(result["repaired_count"], 2)
        self.assertEqual((record.strain, mouse.strain), ("Camk2/Ai93-Ras-N", "Camk2/Ai93-Ras-N"))
        self.assertIsNotNone(self.db.query(Strain).filter_by(name="Camk2/Ai93-Ras-N").first())
        response = list_genotypes(page=1, page_size=50, limit=200, db=self.db, current_user=User())
        self.assertEqual(response["items"][0]["strain"], "Camk2/Ai93-Ras-N")
        self.assertEqual(response["items"][0]["genotype_3"], "RasN-杂合子")
        self.assertEqual(sync_and_normalize_all_strains(self.db)["repaired_count"], 0)

    def test_marker_recognition_accepts_all_orders_but_requires_each_marker(self):
        for genotypes in permutations(("Ai93-杂合子", "Camk2a-阴性", "RasN-野生型")):
            with self.subTest(genotypes=genotypes):
                self.assertEqual(infer_strain_from_genotypes(*genotypes), "Camk2/Ai93-Ras-N")
        for genotypes in [
            ("Ai93-杂合子", "Ai93-野生型", "RasN-杂合子"),
            ("Ai93-杂合子", "Camk2a-阴性", None),
            ("Ai148-杂合子", "Camk2a-阴性", "RasN-野生型"),
            ("Ai930-杂合子", "Camk2a-阴性", "RasN-野生型"),
        ]:
            with self.subTest(genotypes=genotypes):
                self.assertEqual(infer_strain_from_genotypes(*genotypes), "")

    def test_reordered_markers_are_imported_and_backfilled_without_reordering_results(self):
        self.import_rows(
            ["2026-08-18", "E564", None, None, "F", "E392M+E461F、E462F", "Ai93-杂合子", "Camk2a-阴性", "RasN-野生型"],
        )
        mouse = self.db.query(Mouse).filter_by(mouse_code="E564").one()
        record = self.db.query(GenotypeRecord).filter_by(mouse_code="E564").one()
        self.assertEqual((mouse.strain, record.strain), ("Camk2/Ai93-Ras-N", "Camk2/Ai93-Ras-N"))
        mouse.strain = record.strain = ""
        self.db.commit()
        self.assertEqual(sync_and_normalize_all_strains(self.db)["repaired_count"], 2)
        self.assertEqual((mouse.strain, record.strain), ("Camk2/Ai93-Ras-N", "Camk2/Ai93-Ras-N"))
        self.assertEqual((mouse.genotype_1, mouse.genotype_2), ("Ai93-杂合子", "Camk2a-阴性"))
        self.assertEqual((record.genotype_1, record.genotype_2, record.genotype_3),
                         ("Ai93-杂合子", "Camk2a-阴性", "RasN-野生型"))
        self.assertEqual(sync_and_normalize_all_strains(self.db)["repaired_count"], 0)

    def test_legacy_marker_names_distinguish_ras_from_ras_n(self):
        for camk2 in ["Camk2", "Camk2a"]:
            for ras in ["Ras", "RasN", "Ras-N"]:
                for genotypes in permutations(("Ai93-杂合子", f"{camk2}-阴性", f"{ras}-杂合子")):
                    with self.subTest(genotypes=genotypes):
                        expected = "Camk2/Ai93-Ras" if ras == "Ras" else "Camk2/Ai93-Ras-N"
                        self.assertEqual(infer_strain_from_genotypes(*genotypes), expected)
        for genotypes in [
            ("Ai93-杂合子", None, "Ras-杂合子"),
            ("Ai148-杂合子", "Camk2-阳性", "Ras-杂合子"),
            ("Ai93-杂合子", "Camk20-阳性", "Ras-杂合子"),
            ("Ai93-杂合子", "Camk2-阳性", "Rasgrf2-杂合子"),
            ("Ai93-杂合子", "Camk2-阳性", "RasN2-杂合子"),
        ]:
            with self.subTest(genotypes=genotypes):
                self.assertEqual(infer_strain_from_genotypes(*genotypes), "")

    def test_e163_e169_legacy_ras_records_are_imported_and_backfilled(self):
        self.import_rows(*[
            ["2026-01-01", f"E{number}", None, None, "F", "C752M+C940F", "Ai93-杂合子", "Camk2a-阳性", "Ras-杂合子"]
            for number in range(163, 170)
        ])
        for model in [Mouse, GenotypeRecord]:
            for row in self.db.query(model).all():
                self.assertEqual(row.strain, "Camk2/Ai93-Ras")
                row.strain = ""
        self.db.commit()
        self.assertEqual(sync_and_normalize_all_strains(self.db)["repaired_count"], 14)
        for model in [Mouse, GenotypeRecord]:
            for row in self.db.query(model).all():
                self.assertEqual(row.strain, "Camk2/Ai93-Ras")
                self.assertEqual((row.genotype_1, row.genotype_2), ("Ai93-杂合子", "Camk2a-阳性"))
                if isinstance(row, GenotypeRecord):
                    self.assertEqual(row.genotype_3, "Ras-杂合子")
        self.assertEqual(sync_and_normalize_all_strains(self.db)["repaired_count"], 0)

    def test_backfill_preserves_explicit_strains_and_handles_unlinked_record(self):
        mouse = Mouse(mouse_code="F335", strain="自定义品系")
        record = GenotypeRecord(
            mouse_code="F335", strain="", genotype_1="Camk2a-阳性",
            genotype_2="Ai93-杂合子", genotype_3="RasN-杂合子",
        )
        explicit = GenotypeRecord(
            mouse_code="F336", strain="已有品系", genotype_1="Camk2a-阳性",
            genotype_2="Ai93-杂合子", genotype_3="RasN-杂合子",
        )
        ambiguous = GenotypeRecord(mouse_code="F337", strain="", genotype_1="阳性", genotype_2="纯合子", genotype_3="野生型")
        self.db.add_all([mouse, record, explicit, ambiguous])
        self.db.commit()
        sync_and_normalize_all_strains(self.db)
        self.assertEqual((mouse.strain, record.strain), ("自定义品系", "自定义品系"))
        self.assertEqual(explicit.strain, "已有品系")
        self.assertEqual(ambiguous.strain, "")

    def test_th_cre_name_is_imported_without_named_genotype_results(self):
        self.import_rows(
            ["2026-03-01", "F066", "TH-Cre", "2026-02-09", "M", "E640M+E458F、E786F", "阴性", None, None],
        )
        mouse = self.db.query(Mouse).filter_by(mouse_code="F066").one()
        record = self.db.query(GenotypeRecord).filter_by(mouse_code="F066").one()
        self.assertEqual((mouse.strain, record.strain), ("TH-cre", "TH-cre"))
        self.assertEqual(record.genotype_1, "阴性")

    def test_ras_n_name_is_imported_with_correct_capitalization(self):
        self.import_rows(*[
            ["2026-01-01", f"E{number}", "ras-n" if number == 987 else "RAS-N", "2025-12-01", "M", "E780M+E774F", "纯合子", None, None]
            for number in range(987, 993)
        ])
        for number in range(987, 993):
            code = f"E{number}"
            self.assertEqual(self.db.query(Mouse).filter_by(mouse_code=code).one().strain, "RAS-N")
            record = self.db.query(GenotypeRecord).filter_by(mouse_code=code).one()
            self.assertEqual((record.strain, record.genotype_1), ("RAS-N", "纯合子"))

    def test_known_ras_n_mice_are_backfilled_without_changing_their_parents(self):
        for code in ["E780", "E774"]:
            self.db.add(Mouse(mouse_code=code, strain="Ras-N/Ai93"))
        for number in range(987, 993):
            code = f"E{number}"
            mouse = Mouse(mouse_code=code, strain="", parents="E780M+E774F")
            self.db.add(GenotypeRecord(mouse=mouse, mouse_code=code, strain="", parents=mouse.parents, genotype_1="纯合子"))
        self.db.add(Mouse(mouse_code="E9870", strain="", parents="E780M+E774F"))
        self.db.add(Mouse(mouse_code="OTHER1", strain="", parents="E987M+E774F"))
        self.db.commit()

        self.assertEqual(backfill_known_pedigree_strains(self.db), 12)
        for number in range(987, 993):
            code = f"E{number}"
            self.assertEqual(self.db.query(Mouse).filter_by(mouse_code=code).one().strain, "RAS-N")
            self.assertEqual(self.db.query(GenotypeRecord).filter_by(mouse_code=code).one().strain, "RAS-N")
        for code in ["E780", "E774"]:
            self.assertEqual(self.db.query(Mouse).filter_by(mouse_code=code).one().strain, "Ras-N/Ai93")
        for code in ["E9870", "OTHER1"]:
            self.assertEqual(self.db.query(Mouse).filter_by(mouse_code=code).one().strain, "")
        self.assertIsNotNone(self.db.query(Strain).filter_by(name="RAS-N").first())
        self.assertEqual(backfill_known_pedigree_strains(self.db), 0)

    def test_confirmed_ras_n_batches_preserve_results_and_leading_zero_codes(self):
        definitions = [
            (f"E{number:03d}", genotype, "C795M+C790F、C791F")
            for number, genotype in [
                *zip(range(19, 28), ["野生型", "杂合子", "杂合子", "野生型", "野生型", "杂合子", "杂合子", "纯合子", "纯合子"]),
                *zip(range(87, 93), ["野生型", "野生型", "杂合子", "野生型", "杂合子", "纯合子"]),
            ]
        ]
        definitions.extend((f"C{number}", "纯合子", "C796M+C792F、C793F") for number in range(968, 973))
        definitions.extend(
            (f"C{number}", genotype, "C795M+C790F、C791F")
            for number, genotype in zip(range(915, 924), ["纯合子", "野生型", "纯合子", "杂合子", "杂合子", "纯合子", "杂合子", "野生型", "纯合子"])
        )
        unrelated_codes = [
            "E018", "E028", "E0190", "E19", "E086", "E093", "E0870", "E87", "C7950", "C7900", "C7910",
            "C967", "C973", "C9680", "E968", "C7960", "C7920", "C7930",
            "C914", "C924", "C9150", "E915",
        ]
        for code, genotype, parents in definitions:
            mouse = Mouse(mouse_code=code, strain="", parents=parents, genotype_1=genotype)
            self.db.add(GenotypeRecord(
                mouse=mouse, mouse_code=code, strain="", parents=mouse.parents, genotype_1=genotype,
            ))
        for code in unrelated_codes:
            self.db.add(Mouse(mouse_code=code, strain=""))
        explicit = GenotypeRecord(mouse_code="E087", strain="已有品系", genotype_1="野生型")
        self.db.add(explicit)
        self.db.commit()

        self.assertEqual(backfill_known_pedigree_strains(self.db), 58)
        for code, genotype, parents in definitions:
            mouse = self.db.query(Mouse).filter_by(mouse_code=code).one()
            record = self.db.query(GenotypeRecord).filter_by(mouse_code=code, strain="RAS-N").one()
            for row in [mouse, record]:
                self.assertEqual((row.strain, row.genotype_1, row.parents), ("RAS-N", genotype, parents))
        for code in unrelated_codes:
            self.assertEqual(self.db.query(Mouse).filter_by(mouse_code=code).one().strain, "")
        self.assertEqual(explicit.strain, "已有品系")
        self.assertIsNotNone(self.db.query(Strain).filter_by(name="RAS-N").first())
        self.assertEqual(backfill_known_pedigree_strains(self.db), 0)

    def test_ras_import_remains_distinct_from_ras_n(self):
        self.import_rows(
            ["2026-01-01", "C990", "RAS", None, "M", "C627M+C625F", "纯合子", None, None],
            ["2026-01-02", "E987", "ras-n", None, "M", None, "纯合子", None, None],
        )
        for model in [Mouse, GenotypeRecord]:
            self.assertEqual(self.db.query(model).filter_by(mouse_code="C990").one().strain, "Ras")
            self.assertEqual(self.db.query(model).filter_by(mouse_code="E987").one().strain, "RAS-N")

    def test_confirmed_ras_batches_preserve_results_and_exact_codes(self):
        definitions = [
            (f"C{number}", "C627M+C625F", "纯合子")
            for number in [*range(814, 821), *range(990, 998)]
        ]
        definitions.extend((f"C{number}", "C476M+B780F、B781F", "纯合子") for number in range(821, 824))
        definitions.extend(
            (f"C{number}", "C476M+B780F、B781F", genotype)
            for number, genotype in zip(range(712, 716), ["纯合子", "纯合子", "杂合子", "杂合子"])
        )
        definitions.extend(
            (f"C{number}", "C476M+B780F、B781F", genotype)
            for number, genotype in zip(range(669, 675), ["杂合子", "杂合子", "纯合子", "纯合子", "杂合子", "杂合子"])
        )
        definitions.extend(
            (f"B{number}", "不明" if number in [778, 779, 782] else "B782M+B778F、B779F", genotype)
            for number, genotype in zip(range(778, 789), ["杂合子", "杂合子", "杂合子", "杂合子", "杂合子", "杂合子", "纯合子", "纯合子", "杂合子", "纯合子", "杂合子"])
        )
        for code, parents, genotype in definitions:
            mouse = Mouse(mouse_code=code, strain="", parents=parents, genotype_1=genotype)
            self.db.add(GenotypeRecord(
                mouse=mouse, mouse_code=code, strain="", parents=mouse.parents, genotype_1=genotype,
            ))
        unrelated_codes = [
            "C989", "C998", "C9900", "C627", "C625", "E990",
            "C8130", "C824", "C8140", "E814", "C476",
            "C711", "C716", "C7120", "E712",
            "C675", "C6690", "E669",
            "B777", "B789", "B7780", "E778",
        ]
        for code in unrelated_codes:
            self.db.add(Mouse(mouse_code=code, strain="已有品系" if code == "E990" else ""))
        explicit = GenotypeRecord(mouse_code="C990", strain="已有品系", genotype_1="纯合子")
        self.db.add(explicit)
        self.db.commit()

        self.assertEqual(backfill_known_pedigree_strains(self.db), 78)
        for code, parents, genotype in definitions:
            mouse = self.db.query(Mouse).filter_by(mouse_code=code).one()
            record = self.db.query(GenotypeRecord).filter_by(mouse_code=code, strain="Ras").one()
            for row in [mouse, record]:
                self.assertEqual((row.strain, row.genotype_1, row.parents), ("Ras", genotype, parents))
        for code in unrelated_codes:
            self.assertEqual(self.db.query(Mouse).filter_by(mouse_code=code).one().strain, "已有品系" if code == "E990" else "")
        self.assertEqual(explicit.strain, "已有品系")
        self.assertIsNotNone(self.db.query(Strain).filter_by(name="Ras").first())
        self.assertEqual(backfill_known_pedigree_strains(self.db), 0)

    def test_confirmed_th_cre_batches_preserve_results_and_exact_codes(self):
        numbers = [479, 565, 566, *range(664, 669)]
        for number in numbers:
            code = f"C{number}"
            parents = "C288M+C283F" if number == 479 else "C288M+C283F、C479F"
            mouse = Mouse(mouse_code=code, strain="", parents=parents, genotype_1="阳性")
            self.db.add(GenotypeRecord(
                mouse=mouse, mouse_code=code, strain="", parents=mouse.parents, genotype_1="阳性",
            ))
        unrelated_codes = ["C6640", "E664", "C2880", "C2830", "C478", "C480", "C4790", "E479", "C564", "C567", "C5650", "E565"]
        for code in unrelated_codes:
            self.db.add(Mouse(mouse_code=code, strain=""))
        explicit = GenotypeRecord(mouse_code="C664", strain="TH-cre/Ai148", genotype_1="阳性")
        self.db.add(explicit)
        self.db.commit()

        self.assertEqual(backfill_known_pedigree_strains(self.db), 16)
        for number in numbers:
            code = f"C{number}"
            parents = "C288M+C283F" if number == 479 else "C288M+C283F、C479F"
            mouse = self.db.query(Mouse).filter_by(mouse_code=code).one()
            record = self.db.query(GenotypeRecord).filter_by(mouse_code=code, strain="TH-cre").one()
            for row in [mouse, record]:
                self.assertEqual((row.strain, row.genotype_1, row.parents), ("TH-cre", "阳性", parents))
        for code in unrelated_codes:
            self.assertEqual(self.db.query(Mouse).filter_by(mouse_code=code).one().strain, "")
        self.assertEqual(explicit.strain, "TH-cre/Ai148")
        self.assertEqual(backfill_known_pedigree_strains(self.db), 0)

    def test_confirmed_e798_strain_preserves_genotypes_and_matches_exact_code(self):
        mouse = Mouse(mouse_code="E798", strain="", genotype_2="Ai148-杂合子")
        record = GenotypeRecord(
            mouse=mouse, mouse_code="E798", strain="", genotype_1="Camk2a-阳性",
            genotype_2="Ai148-杂合子", genotype_3="RasN-野生型",
        )
        explicit = GenotypeRecord(mouse_code="E798", strain="已有品系", genotype_2="Ai148-杂合子")
        similar = Mouse(mouse_code="E7980", strain="", genotype_2="Ai148-杂合子")
        self.db.add_all([record, explicit, similar])
        self.db.commit()

        self.assertEqual(sync_and_normalize_all_strains(self.db)["repaired_count"], 2)
        self.assertEqual((mouse.strain, record.strain), ("Camk2/Ai93-Ras-N", "Camk2/Ai93-Ras-N"))
        self.assertEqual((mouse.genotype_2, record.genotype_2), ("Ai148-杂合子", "Ai148-杂合子"))
        self.assertEqual((record.genotype_1, record.genotype_3), ("Camk2a-阳性", "RasN-野生型"))
        self.assertEqual(explicit.strain, "已有品系")
        self.assertEqual(similar.strain, "")
        self.assertEqual(sync_and_normalize_all_strains(self.db)["repaired_count"], 0)

    def test_confirmed_family_follows_exact_pedigrees_and_stops_at_crosses(self):
        definitions = [
            ("C663", "", "C288M+C283F"),
            ("C786", "", "C663M+C664F"),
            ("E458", "", "C786M+C787F"),
            ("F066", "", "E640M+E458F、E786F"),
            ("C288", "", None),
            ("C664", "", None),
            ("OTHER1", "", "C6630M+C664F"),
            ("OTHER2", "", "C663_2M+C664F"),
            ("CROSS1", "TH-cre/Ai148", "C663M+C664F"),
            ("CROSS2", "", "CROSS1M+C664F"),
        ]
        for code, strain, parents in definitions:
            mouse = Mouse(mouse_code=code, strain=strain, parents=parents)
            self.db.add(GenotypeRecord(mouse=mouse, mouse_code=code, strain=strain, parents=parents, genotype_1="阴性"))
        self.db.add(GenotypeRecord(mouse_code="ARCHIVE1", strain="", parents="c663f，C664M", genotype_1="阳性"))
        self.db.commit()

        result = sync_and_normalize_all_strains(self.db)
        self.assertEqual(result["repaired_count"], 13)
        for code in ["C288", "C663", "C664", "C786", "E458", "F066"]:
            self.assertEqual(self.db.query(Mouse).filter_by(mouse_code=code).one().strain, "TH-cre")
            record = self.db.query(GenotypeRecord).filter_by(mouse_code=code).one()
            self.assertEqual((record.strain, record.genotype_1), ("TH-cre", "阴性"))
        for code in ["OTHER1", "OTHER2", "CROSS2"]:
            self.assertEqual(self.db.query(Mouse).filter_by(mouse_code=code).one().strain, "")
        self.assertEqual(self.db.query(Mouse).filter_by(mouse_code="CROSS1").one().strain, "TH-cre/Ai148")
        self.assertEqual(self.db.query(GenotypeRecord).filter_by(mouse_code="ARCHIVE1").one().strain, "TH-cre")
        self.assertIsNone(self.db.query(Mouse).filter_by(mouse_code="ARCHIVE1").first())
        self.assertEqual(backfill_known_pedigree_strains(self.db), 0)

    def test_uploaded_workbook_repairs_family_without_creating_missing_founder(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "基因鉴定结果"
        sheet.append(["测试日期", "小鼠编号", "品系", "父母", "Genotype 1"])
        sheet.append(["2026-03-01", "C786", None, "C663M+C664F", "阴性"])
        with patch("backend.app.services.importer.openpyxl.load_workbook", return_value=workbook):
            result = import_single_excel_file(self.db, "test.xlsx")
        self.assertEqual(result["errors"], [])
        self.assertEqual(result["strains_repaired"], 2)
        self.assertEqual(self.db.query(Mouse).filter_by(mouse_code="C786").one().strain, "TH-cre")
        self.assertIsNone(self.db.query(Mouse).filter_by(mouse_code="C663").first())


if __name__ == "__main__":
    unittest.main()
