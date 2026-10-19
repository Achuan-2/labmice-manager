import io
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from backend.app.auth import require_admin, require_auth
from backend.app.database import Base, get_db
from backend.app.models.models import Mouse, TransferLog, User
from backend.app.routers.import_export import router
from backend.app.routers.transfers import router as transfers_router
from backend.app.routers.genotypes import router as genotypes_router


class ImportReportApiTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.folder = tempfile.TemporaryDirectory()
        self.paths = patch("backend.app.routers.import_export.BASE_DIR", self.folder.name)
        self.paths.start()
        app = FastAPI()
        app.include_router(router)
        app.include_router(transfers_router)
        app.include_router(genotypes_router)

        def test_db():
            with Session(self.engine, autoflush=False) as db:
                yield db

        app.dependency_overrides[get_db] = test_db
        app.dependency_overrides[require_admin] = lambda: User(username="test", display_name="测试管理员", role="admin")
        app.dependency_overrides[require_auth] = app.dependency_overrides[require_admin]
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        self.engine.dispose()
        self.paths.stop()
        self.folder.cleanup()

    @staticmethod
    def excel_content():
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "小鼠档案"
        sheet.append(["日期", "品系", "DOB", "性别", "父母", "编号"])
        sheet.append(["2026-10-01", "测试品系", None, "M", None, "D001"])
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def test_batch_upload_merges_added_and_skipped_reports(self):
        content = self.excel_content()
        response = self.client.post("/api/import-export/upload", files=[
            ("files", ("first.xlsx", content)), ("files", ("second.xlsx", content)),
        ])
        self.assertEqual(response.status_code, 200)
        result = response.json()
        self.assertEqual(result["details"]["mice_imported"], 1)
        self.assertEqual(result["details"]["cages_imported"], 0)
        self.assertEqual(result["file_summaries"][1]["imported_count"], 0)
        mice = [row for row in result["details"]["report"]["records"] if row["type"] == "小鼠"]
        self.assertEqual([(row["status"], row["filename"]) for row in mice], [("added", "first.xlsx"), ("skipped", "second.xlsx")])
        self.assertIn("重复跳过 2 项", result["message"])
        repeat = self.client.post("/api/import-export/upload", files={"files": ("first.xlsx", content)}).json()
        self.assertEqual(repeat["imported_count"], 0)
        self.assertEqual(repeat["details"]["report"]["summary"]["skipped"], 2)
        logs = self.client.get("/api/transfers", params={"action_type": "导入数据"}).json()
        self.assertEqual(logs["total"], 1)
        log = logs["items"][0]
        self.assertEqual(log["operator"], "测试管理员")
        self.assertEqual((log["mouse_codes"], log["mouse_count"]), ("D001", 1))
        self.assertIn("first.xlsx", log["notes"])
        self.assertIn("小鼠档案 · 第2行", log["notes"])
        self.assertIn("【新增】小鼠：D001", log["notes"])

    def test_invalid_workbook_returns_failure_details_without_modifying_data(self):
        with Session(self.engine) as db:
            db.add(Mouse(mouse_code="D001", status="死亡"))
            db.commit()
        response = self.client.post("/api/import-export/upload", files={"files": ("broken.xlsx", b"not an Excel workbook")})
        self.assertEqual(response.status_code, 200)
        result = response.json()
        self.assertEqual(result["imported_count"], 0)
        self.assertEqual(len(result["details"]["errors"]), 1)
        self.assertIn("broken.xlsx", result["details"]["errors"][0])
        with Session(self.engine) as db:
            self.assertEqual(db.query(Mouse).one().status, "死亡")
            self.assertEqual(db.query(TransferLog).count(), 0)

    def test_archive_import_is_visible_in_genotype_history_api(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "小鼠档案"
        sheet.append(["测试日期", "品系", "DOB", "性别", "父母", "编号", "Genotype 1", "Genotype 2", "操作记录", "备注"])
        sheet.append(["/", "Pemt-KO CAS9", "2026-09-18", "M", "H016M+H018F", "H515", "纯合子", None, None, "AJ5"])
        stream = io.BytesIO()
        workbook.save(stream)
        content = stream.getvalue()
        response = self.client.post("/api/import-export/upload", files={"files": ("archive.xlsx", content)})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["details"]["genotypes_imported"], 1)
        history = self.client.get("/api/genotypes/by-mouse/H515")
        self.assertEqual(history.status_code, 200)
        records = history.json()
        self.assertEqual(len(records), 1)
        self.assertEqual((records[0]["genotype_1"], records[0]["test_date"]), ("纯合子", None))
        self.assertIsNotNone(records[0]["mouse_id"])
        listing = self.client.get("/api/genotypes", params={"mouse_code": "H515", "page": 1}).json()
        self.assertEqual(listing["total"], 1)
        repeat = self.client.post("/api/import-export/upload", files={"files": ("archive.xlsx", content)}).json()
        self.assertEqual(repeat["details"]["genotypes_imported"], 0)

    def test_excel_row_order_is_reversed_by_default_even_with_mixed_dates_and_blank_results(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "小鼠档案"
        sheet.append(["测试日期", "品系", "DOB", "性别", "父母", "编号", "Genotype 1", "Genotype 2"])
        for code, date, result in (("H489", "2026-10-18", "WT"), ("H490", "2026-09-01", "HET"), ("H491", None, None)):
            sheet.append([date, "5xFAD", None, "F", None, code, result])
        stream = io.BytesIO()
        workbook.save(stream)
        response = self.client.post("/api/import-export/upload", files={"files": ("ordered.xlsx", stream.getvalue())})
        self.assertEqual(response.status_code, 200)
        default = self.client.get('/api/genotypes', params={"page": 1}).json()
        by_date = self.client.get('/api/genotypes', params={"page": 1, "sort_by": "test_date"}).json()
        self.assertEqual([r["mouse_code"] for r in default["items"]], ["H491", "H490", "H489"])
        self.assertEqual([r["mouse_code"] for r in by_date["items"]], ["H489", "H490", "H491"])


if __name__ == "__main__":
    unittest.main()
