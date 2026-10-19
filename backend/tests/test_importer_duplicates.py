import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from backend.app.database import Base
from backend.app.models.models import Cage, GenotypeRecord, Mouse, TransferLog, TransferRequest
from backend.app.services.importer import import_single_excel_file, normalize_strain_name


class ImportDuplicateTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        # Production sessions disable autoflush; tests must exercise that behavior.
        self.db = Session(self.engine, autoflush=False)
        self.folder = tempfile.TemporaryDirectory()
        self.path = Path(self.folder.name) / "test.xlsx"

    def tearDown(self):
        self.db.close()
        self.engine.dispose()
        self.folder.cleanup()

    def import_book(self, workbook):
        workbook.save(self.path)
        return import_single_excel_file(self.db, str(self.path))

    @staticmethod
    def archive_book(*codes):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "小鼠档案"
        sheet.append(["日期", "品系", "DOB", "性别", "父母", "编号", "Genotype 1", "Genotype 2", "操作", "备注"])
        for code in codes:
            sheet.append(["2026-10-01", "测试品系", "2026-09-01", "M", None, code])
        return workbook

    def test_duplicate_rows_and_incremental_file_only_add_new_codes(self):
        workbook = self.archive_book("D001", "D001")
        first = self.import_book(workbook)
        self.assertEqual(first["errors"], [])
        self.assertEqual(first["mice_imported"], 1)
        second = self.import_book(workbook)
        self.assertEqual(second["mice_imported"], 0)
        self.assertEqual(second["report"]["summary"]["added"], 0)
        self.assertEqual(second["report"]["summary"]["updated"], 0)
        workbook.active.append(["2026-10-02", "测试品系", None, "F", None, "D002"])
        third = self.import_book(workbook)
        self.assertEqual(third["mice_imported"], 1)
        added = [item for item in third["report"]["records"] if item["status"] == "added" and item["type"] == "小鼠"]
        self.assertEqual([(item["identifier"], item["sheet"], item["row"]) for item in added], [("D002", "小鼠档案", 4)])
        self.assertEqual(self.db.query(Mouse).count(), 2)

    def test_reimport_preserves_manually_changed_cage_and_status(self):
        workbook = self.archive_book("D001")
        sheet = workbook.create_sheet("小鼠笼位信息")
        sheet.append(["笼位", "品系", "数量", "性别", "编号"])
        sheet.append(["A1", "测试品系", 1, "M", "D001"])
        self.import_book(workbook)
        mouse = self.db.query(Mouse).one()
        new_cage = Cage(room="测试鼠房", cage_code="B1")
        self.db.add(new_cage)
        self.db.flush()
        mouse.cage_id = new_cage.id
        mouse.status = "自定义状态"
        mouse.owner_name = "管理员设置的领取人"
        self.db.commit()
        report = self.import_book(workbook)
        self.assertEqual(report["report"]["summary"]["updated"], 0)
        self.assertEqual((mouse.cage_id, mouse.status, mouse.owner_name), (new_cage.id, "自定义状态", "管理员设置的领取人"))
        mouse.cage_id = None
        mouse.status = "死亡"
        self.db.commit()
        self.import_book(workbook)
        self.assertIsNone(mouse.cage_id)
        self.assertEqual(mouse.status, "死亡")

    def test_generated_codes_are_stable_even_after_manually_moving_mouse(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "小鼠笼位信息"
        sheet.append(["笼位", "品系", "数量", "性别", "编号"])
        sheet.append(["A1", "测试品系", 1, "M", "无耳标"])
        first = self.import_book(workbook)
        mouse = self.db.query(Mouse).one()
        mouse.cage_id = None
        self.db.commit()
        second = self.import_book(workbook)
        self.assertEqual(first["mice_imported"], 1)
        self.assertEqual(second["errors"], [])
        self.assertEqual(second["mice_imported"], 0)
        self.assertEqual(self.db.query(Mouse).count(), 1)
        self.assertIsNone(mouse.cage_id)

    def test_transfer_requests_and_logs_do_not_repeat(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "转鼠申请"
        sheet.append(["转鼠申请"])
        sheet.append(["序号", "日期", "领取人", "品系", "要求", "转入", "笼数", "转出", "性别", "编号"])
        row = ["1", "2026-10-01", "测试领取人", "测试品系", "成年", "东五", 1, "江湾", "M", "D001"]
        sheet.append(row)
        sheet.append(row)
        self.import_book(workbook)
        second = self.import_book(workbook)
        self.assertEqual(second["errors"], [])
        self.assertEqual((second["transfer_requests_imported"], second["transfers_imported"]), (0, 0))
        self.assertEqual(self.db.query(TransferRequest).count(), 1)
        self.assertEqual(self.db.query(TransferLog).filter(TransferLog.action_type == "转鼠/领用").count(), 1)

    def test_incremental_import_logs_only_changed_codes(self):
        workbook = self.archive_book("D001")
        self.import_book(workbook)
        workbook.active.append(["2026-10-02", "测试品系", None, "F", None, "D002"])
        second = self.import_book(workbook)
        logs = self.db.query(TransferLog).filter_by(action_type="导入数据").order_by(TransferLog.id).all()
        self.assertEqual(len(logs), 2)
        self.assertEqual(logs[1].mouse_codes, "D002")
        self.assertIn("【新增】小鼠：D002", logs[1].notes)
        self.assertNotIn("【新增】小鼠：D001", logs[1].notes)
        # The audit entry itself is not a newly imported transfer record.
        self.assertEqual(second["transfers_imported"], 0)
        self.import_book(workbook)
        self.assertEqual(self.db.query(TransferLog).filter_by(action_type="导入数据").count(), 2)

    def test_import_and_its_log_rollback_together_if_logging_fails(self):
        with patch("backend.app.services.importer.record_import_logs", side_effect=RuntimeError("测试日志失败")):
            with self.assertRaisesRegex(RuntimeError, "测试日志失败"):
                self.import_book(self.archive_book("D001"))
        self.assertEqual(self.db.query(Mouse).count(), 0)
        self.assertEqual(self.db.query(TransferLog).count(), 0)

    def test_update_report_contains_only_changed_fields(self):
        self.db.add(Mouse(mouse_code="D001", strain="", status="出笼"))
        self.db.commit()
        result = self.import_book(self.archive_book("D001"))
        self.assertEqual(result["mice_imported"], 0)
        updated = [item for item in result["report"]["records"] if item["status"] == "updated"]
        self.assertEqual(len(updated), 1)
        self.assertIn({"field": "品系", "before": "", "after": "测试品系"}, updated[0]["changes"])
        log = self.db.query(TransferLog).filter_by(action_type="导入数据").one()
        self.assertIn("【补充/更新】小鼠：D001", log.notes)
        self.assertIn("品系：空 → 测试品系", log.notes)

    def test_bad_row_does_not_rollback_prior_or_following_rows(self):
        workbook = self.archive_book("D001", "D002", "D003")
        workbook.active.cell(3, 2, "坏数据")

        def normalize(db, value):
            if value == "坏数据":
                raise ValueError("测试坏行")
            return normalize_strain_name(db, value)

        with patch("backend.app.services.importer.normalize_strain_name", side_effect=normalize):
            result = self.import_book(workbook)
        self.assertEqual(self.db.query(Mouse).count(), 2)
        self.assertEqual(result["mice_imported"], 2)
        self.assertEqual(result["errors"], ["小鼠档案 第3行解析错误: 测试坏行"])

    def test_third_genotype_distinguishes_new_test_results(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "基因鉴定结果"
        sheet.append(["日期", "编号", "品系", "Genotype 1", "Genotype 2", "Genotype 3"])
        sheet.append(["2026-10-01", "D001", "测试品系", "HET", "WT", "阳性"])
        sheet.append(["2026-10-01", "D001", "测试品系", "HET", "WT", "阴性"])
        first = self.import_book(workbook)
        second = self.import_book(workbook)
        self.assertEqual(first["genotypes_imported"], 2)
        self.assertEqual(second["genotypes_imported"], 0)
        self.assertEqual(self.db.query(GenotypeRecord).count(), 2)

    def test_archive_result_without_test_date_backfills_h515_history(self):
        mouse = Mouse(
            mouse_code="H515", strain="Pemt-KO CAS9", gender="M", dob="2026-09-18",
            parents="H016M+H018F", genotype_1="纯合子", genotype_2="", source_room="江湾发育所",
        )
        self.db.add(mouse)
        self.db.commit()
        workbook = self.archive_book()
        workbook.active.append(["/", "Pemt-KO CAS9", "2026-09-18", "M", "H016M+H018F", "H515", "纯合子", None, None, "AJ5"])
        first = self.import_book(workbook)
        record = self.db.query(GenotypeRecord).filter_by(mouse_code="H515").one()
        self.assertEqual((record.mouse_id, record.genotype_1, record.test_date), (mouse.id, "纯合子", None))
        self.assertEqual((record.dob, record.parents, record.notes), ("2026-09-18", "H016M+H018F", "AJ5"))
        self.assertEqual((first["mice_imported"], first["genotypes_imported"]), (0, 1))
        self.assertEqual(mouse.genotype_2, "")
        self.assertFalse(any(item["type"] == "小鼠" and item["status"] == "updated" for item in first["report"]["records"]))
        added = [item for item in first["report"]["records"] if item["type"] == "基因鉴定"]
        self.assertEqual((added[0]["status"], added[0]["sheet"], added[0]["row"]), ("added", "小鼠档案", 2))
        log = self.db.query(TransferLog).filter_by(action_type="导入数据").one()
        self.assertIn("【新增】基因鉴定：H515 / 无日期 / 纯合子", log.notes)
        second = self.import_book(workbook)
        self.assertEqual(second["report"]["summary"]["added"], 0)
        self.assertEqual(second["report"]["summary"]["updated"], 0)
        self.assertEqual(self.db.query(GenotypeRecord).count(), 1)
        self.assertEqual(self.db.query(TransferLog).filter_by(action_type="导入数据").count(), 1)

    def test_archive_blank_results_are_registered_and_reimported_without_duplicates(self):
        workbook = self.archive_book()
        workbook.active.append([None, "5xFAD", "2026-09-18", "F", "H087M+H091F、H092F", "H489", None, None, None, "AI5"])
        workbook.active.append(["/", "fos-GFP", "2026-09-18", "M", "D561M+D557F", "H494", "/", " ", None, "AI7"])
        first = self.import_book(workbook)
        self.assertEqual(first["errors"], [])
        self.assertEqual(first["genotypes_imported"], 2)
        records = self.db.query(GenotypeRecord).order_by(GenotypeRecord.mouse_code).all()
        self.assertEqual([r.mouse_code for r in records], ["H489", "H494"])
        self.assertTrue(all(r.genotype_1 is None and r.genotype_2 is None and r.test_date is None for r in records))
        self.assertEqual((records[0].dob, records[0].gender, records[0].parents, records[0].notes),
                         ("2026-09-18", "F", "H087M+H091F、H092F", "AI5"))
        self.assertTrue(all(r.mouse_id is not None for r in records))
        second = self.import_book(workbook)
        self.assertEqual(second["report"]["summary"]["added"], 0)
        self.assertEqual(second["report"]["summary"]["updated"], 0)
        self.assertEqual(self.db.query(GenotypeRecord).count(), 2)

    def test_blank_archive_does_not_clear_manually_recorded_results(self):
        mouse = Mouse(mouse_code="H489", genotype_1="HET", genotype_2="WT", test_date="2026-10-01")
        self.db.add(mouse)
        self.db.commit()
        self.import_book(self.archive_book("H489"))
        self.assertEqual((mouse.genotype_1, mouse.genotype_2, mouse.test_date), ("HET", "WT", "2026-10-01"))
        self.assertIsNone(self.db.query(GenotypeRecord).one().genotype_1)

    def test_archive_and_genotype_sheets_share_empty_value_deduplication(self):
        mouse = Mouse(mouse_code="H515", source_room="江湾发育所")
        record = GenotypeRecord(
            mouse_code="H515", test_date="", genotype_1="纯合子", genotype_2="", genotype_3="/",
            notes="人工备注", op_record="人工录入",
        )
        self.db.add_all([mouse, record])
        self.db.commit()
        existing_id = record.id
        workbook = self.archive_book()
        workbook.active.append(["/", "测试品系", None, "M", None, "H515", "纯合子", "/", "Excel操作", "Excel备注"])
        sheet = workbook.create_sheet("基因鉴定结果")
        sheet.append(["日期", "编号", "品系", "Genotype 1", "Genotype 2", "Genotype 3"])
        sheet.append([None, "H515", "测试品系", "纯合子", None, None])
        result = self.import_book(workbook)
        self.assertEqual(result["errors"], [])
        self.assertEqual(result["genotypes_imported"], 0)
        self.assertEqual(self.db.query(GenotypeRecord).count(), 1)
        self.assertEqual((record.id, record.mouse_id, record.strain), (existing_id, mouse.id, "测试品系"))
        self.assertEqual((record.notes, record.op_record), ("人工备注", "人工录入"))

    def test_archive_history_syncs_latest_summary_and_reads_third_outcome(self):
        workbook = self.archive_book()
        sheet = workbook.active
        sheet.cell(1, 11, "Genotype 3")
        sheet.append(["2026-10-02", "测试品系", None, "M", None, "D001", "阳性", None, None, None, "HET"])
        sheet.append(["2026-09-01", "测试品系", None, "M", None, "D001", "阴性"])
        sheet.append([None, "测试品系", None, "M", None, "D001", "待测"])
        result = self.import_book(workbook)
        self.assertEqual(result["genotypes_imported"], 3)
        mouse = self.db.query(Mouse).one()
        self.assertEqual((mouse.genotype_1, mouse.test_date), ("阳性", "2026-10-02"))
        record = self.db.query(GenotypeRecord).filter_by(test_date="2026-10-02").one()
        self.assertEqual(record.genotype_3, "HET")


if __name__ == "__main__":
    unittest.main()
